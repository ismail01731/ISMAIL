from typing import Optional
import base64
import hashlib
import hmac
import json
import os
import io
import csv
import re
import time
import secrets

from dotenv import load_dotenv

from fastapi import (
    FastAPI,
    HTTPException,
    Request,
    UploadFile,
    File,
    Depends,
    WebSocket,
    WebSocketDisconnect,
)
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials



from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

load_dotenv(
    os.path.join(
        os.path.dirname(__file__),
        "..",
        "config",
        ".env"
    ),
    override=False
)


from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from backend.ai_engine import AIEngine
from backend.input_security import InputSecurity
from backend.web_research import WebResearch
from backend.knowledge_base import KnowledgeBase








app = FastAPI(
    title="ISMAIL AI",
    version="0.1.0",
    description="ISMAIL AI backend"
)

security = HTTPBearer()

FRONTEND_DIR = os.path.abspath(
    os.path.join(os.path.dirname(__file__), "..", "frontend")
)

if os.path.isdir(FRONTEND_DIR):
    app.mount("/app", StaticFiles(directory=FRONTEND_DIR, html=True), name="frontend")


app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


ai_engine = AIEngine()
web_research = WebResearch()
knowledge_base = KnowledgeBase()
input_security = InputSecurity()


AUTH_SECRET = os.getenv("AUTH_SECRET", "").strip()
ADMIN_KNOWLEDGE_KEY = os.getenv(
    "ADMIN_KNOWLEDGE_KEY",
    "",
).strip()

IDENTITY_TOKEN_MAX_AGE_SECONDS = 30 * 24 * 60 * 60
IDENTITY_TOKEN_CLOCK_SKEW_SECONDS = 60

if not AUTH_SECRET:
    raise RuntimeError(
        "AUTH_SECRET is required for authentication."
    )

if not ADMIN_KNOWLEDGE_KEY:
    raise RuntimeError(
        "ADMIN_KNOWLEDGE_KEY is required for knowledge administration."
    )




def _generate_identity() -> str:
    return "user-" + secrets.token_urlsafe(24)


def _hash_password(password: str) -> str:
    password_bytes = password.encode("utf-8")
    salt = secrets.token_bytes(16)
    iterations = 310000

    password_hash = hashlib.pbkdf2_hmac(
        "sha256",
        password_bytes,
        salt,
        iterations,
    )

    encoded_salt = base64.urlsafe_b64encode(
        salt
    ).decode("ascii").rstrip("=")

    encoded_hash = base64.urlsafe_b64encode(
        password_hash
    ).decode("ascii").rstrip("=")

    return f"pbkdf2_sha256${iterations}${encoded_salt}${encoded_hash}"


def _verify_password(
    password: str,
    stored_hash: str,
) -> bool:
    try:
        algorithm, iterations_text, encoded_salt, encoded_hash = (
            stored_hash.split("$", 3)
        )

        if algorithm != "pbkdf2_sha256":
            return False

        iterations = int(iterations_text)

        salt = base64.urlsafe_b64decode(
            encoded_salt + "=" * (-len(encoded_salt) % 4)
        )

        expected_hash = base64.urlsafe_b64decode(
            encoded_hash + "=" * (-len(encoded_hash) % 4)
        )

        actual_hash = hashlib.pbkdf2_hmac(
            "sha256",
            password.encode("utf-8"),
            salt,
            iterations,
        )

        return hmac.compare_digest(
            actual_hash,
            expected_hash,
        )

    except (ValueError, TypeError):
        return False




def _sign_identity(user_id: str, issued_at: int) -> str:
    payload = json.dumps(
        {
            "user_id": user_id,
            "iat": issued_at,
        },
        separators=(",", ":"),
        sort_keys=True,
    ).encode("utf-8")
    encoded_payload = base64.urlsafe_b64encode(payload).decode("ascii").rstrip("=")
    signature = hmac.new(
        AUTH_SECRET.encode("utf-8"),
        encoded_payload.encode("ascii"),
        hashlib.sha256,
    ).digest()
    encoded_signature = base64.urlsafe_b64encode(signature).decode("ascii").rstrip("=")
    return f"{encoded_payload}.{encoded_signature}"

def _verify_identity(token: str) -> str:
    if not isinstance(token, str) or token.count(".") != 1:
        raise ValueError("Invalid identity token.")
    encoded_payload, encoded_signature = token.split(".", 1)
    try:
        payload_bytes = base64.urlsafe_b64decode(
            encoded_payload + "=" * (-len(encoded_payload) % 4)
        )
        provided_signature = base64.urlsafe_b64decode(
            encoded_signature + "=" * (-len(encoded_signature) % 4)
        )
        payload = json.loads(payload_bytes.decode("utf-8"))
    except (ValueError, TypeError, json.JSONDecodeError):
        raise ValueError("Invalid identity token.")
    expected_signature = hmac.new(
        AUTH_SECRET.encode("utf-8"),
        encoded_payload.encode("ascii"),
        hashlib.sha256,
    ).digest()
    if not hmac.compare_digest(provided_signature, expected_signature):
        raise ValueError("Invalid identity token.")
    user_id = payload.get("user_id")
    issued_at = payload.get("iat")
    if not isinstance(user_id, str) or not user_id:
        raise ValueError("Invalid identity token.")
    if not isinstance(issued_at, int) or isinstance(issued_at, bool):
        raise ValueError("Invalid identity token.")
    now = int(time.time())
    if issued_at > now + IDENTITY_TOKEN_CLOCK_SKEW_SECONDS:
        raise ValueError("Identity token is not yet valid.")
    if now - issued_at > IDENTITY_TOKEN_MAX_AGE_SECONDS:
        raise ValueError("Identity token has expired.")
    return user_id

class ChatRateLimiter:
    """Simple in-memory per-IP rate limiter."""

    def __init__(
        self,
        max_requests: int,
        window_seconds: int,
    ):
        self.max_requests = max_requests
        self.window_seconds = window_seconds
        self._requests = {}

    def check(self, client_ip: str) -> bool:
        now = time.monotonic()
        timestamps = self._requests.get(client_ip, [])

        timestamps = [
            timestamp
            for timestamp in timestamps
            if now - timestamp < self.window_seconds
        ]

        if len(timestamps) >= self.max_requests:
            self._requests[client_ip] = timestamps
            return False

        timestamps.append(now)
        self._requests[client_ip] = timestamps
        return True


chat_rate_limiter = ChatRateLimiter(
    max_requests=10,
    window_seconds=60,
)

research_rate_limiter = ChatRateLimiter(
    max_requests=3,
    window_seconds=60,
)


auth_register_rate_limiter = ChatRateLimiter(
    max_requests=5,
    window_seconds=60,
)

auth_login_rate_limiter = ChatRateLimiter(
    max_requests=5,
    window_seconds=60,
)

session_rate_limiter = ChatRateLimiter(
    max_requests=5,
    window_seconds=60,
)


def validate_registration_password(password: str):
    if len(password) < 8:
        return "Password must contain at least 8 characters."

    if len(password) > 128:
        return "Password must not exceed 128 characters."

    if not re.search(r"[a-z]", password):
        return "Password must contain at least one lowercase letter."

    if not re.search(r"[A-Z]", password):
        return "Password must contain at least one uppercase letter."

    if not re.search(r"[0-9]", password):
        return "Password must contain at least one number."

    if not re.search(r"[^A-Za-z0-9]", password):
        return "Password must contain at least one special character."

    return None





class RegisterRequest(BaseModel):
    username: str = Field(
        ...,
        min_length=3,
        max_length=50,
        pattern=r"^[A-Za-z0-9._-]+$",
    )
    password: str = Field(
        ...,
        min_length=8,
        max_length=128,
    )


class LoginRequest(BaseModel):
    username: str = Field(
        ...,
        min_length=3,
        max_length=50,
        pattern=r"^[A-Za-z0-9._-]+$",
    )
    password: str = Field(
        ...,
        min_length=8,
        max_length=128,
    )




@app.post("/api/auth/register")
def register_account(
    request: RegisterRequest,
    http_request: Request,
):
    try:
        client_ip = (
            http_request.client.host
            if http_request.client
            else "unknown"
        )

        if not auth_register_rate_limiter.check(client_ip):
            raise HTTPException(
                status_code=429,
                detail=(
                    "Too many registration requests. "
                    "Please try again later."
                ),
            )
        username = request.username.strip().lower()


        password_error = validate_registration_password(
            request.password
        )

        if password_error:
            raise HTTPException(
                status_code=400,
                detail=password_error,
            )

        if len(username) < 3:
            raise HTTPException(
                status_code=400,
                detail="Username must contain at least 3 characters.",
            )

        user_id = _generate_identity()

        password_hash = _hash_password(request.password)

        created = knowledge_base.create_user_account(
            user_id,
            username,
            password_hash,
        )

        if not created:
            raise HTTPException(
                status_code=409,
                detail="Username is already registered.",
            )

        issued_at = int(time.time())
        token = _sign_identity(
            user_id,
            issued_at,
        )

        return {
            "name": "ISMAIL AI",
            "user_id": user_id,
            "username": username,
            "token": token,
            "expires_in": IDENTITY_TOKEN_MAX_AGE_SECONDS,
        }

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail="Registration failed. Please try again later.",
        )


@app.post("/api/auth/login")
def login_account(
    request: LoginRequest,
    http_request: Request,
):
    try:
        client_ip = (
            http_request.client.host
            if http_request.client
            else "unknown"
        )

        if not auth_login_rate_limiter.check(client_ip):
            raise HTTPException(
                status_code=429,
                detail=(
                    "Too many login requests. "
                    "Please try again later."
                ),
            )
        username = request.username.strip().lower()

        account = knowledge_base.get_user_account(
            username
        )

        if account is None:
            raise HTTPException(
                status_code=401,
                detail="Invalid username or password.",
            )

        password_ok = _verify_password(
            request.password,
            account["password_hash"],
        )

        if not password_ok:
            raise HTTPException(
                status_code=401,
                detail="Invalid username or password.",
            )

        issued_at = int(time.time())

        token = _sign_identity(
            account["user_id"],
            issued_at,
        )

        return {
            "name": "ISMAIL AI",
            "user_id": account["user_id"],
            "username": account["username"],
            "token": token,
            "expires_in": IDENTITY_TOKEN_MAX_AGE_SECONDS,
        }

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail="Login failed. Please try again later.",
        )


class GoogleLoginRequest(BaseModel):
    credential: str = Field(..., min_length=10)


@app.post("/api/auth/google")
def google_login(
    request: GoogleLoginRequest,
    http_request: Request,
):
    try:
        from google.oauth2 import id_token
        from google.auth.transport import requests as google_requests

        google_user = id_token.verify_oauth2_token(
            request.credential,
            google_requests.Request(),
            os.getenv("GOOGLE_CLIENT_ID", "").strip(),
        )

        google_email = google_user.get("email", "").strip().lower()

        if not google_email:
            raise HTTPException(
                status_code=401,
                detail="Google account email পাওয়া যায়নি.",
            )

        user_id = "google-" + hashlib.sha256(
            google_email.encode("utf-8")
        ).hexdigest()[:24]

        username = google_email.split("@")[0]

        issued_at = int(time.time())

        token = _sign_identity(
            user_id,
            issued_at,
        )

        return {
            "name": "ISMAIL AI",
            "user_id": user_id,
            "username": username,
            "token": token,
            "expires_in": IDENTITY_TOKEN_MAX_AGE_SECONDS,
            "email": google_email,
        }

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=401,
            detail="Google Login failed. Please try again later.",
        )


@app.post("/api/session")
def create_session(http_request: Request):
    client_ip = (
        http_request.client.host
        if http_request.client
        else "unknown"
    )

    if not session_rate_limiter.check(client_ip):
        raise HTTPException(
            status_code=429,
            detail=(
                "Too many session requests. "
                "Please try again later."
            ),
        )

    user_id = _generate_identity()
    issued_at = int(time.time())
    token = _sign_identity(user_id, issued_at)
    return {
        "user_id": user_id,
        "token": token,
        "expires_in": IDENTITY_TOKEN_MAX_AGE_SECONDS,
    }




# =========================================================
# FILE UPLOAD AND TEXT EXTRACTION
# =========================================================

MAX_UPLOAD_SIZE = 10 * 1024 * 1024
MAX_EXTRACTED_TEXT = 50000


def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
    name = filename.lower()

    if name.endswith(".txt") or name.endswith(".md") or name.endswith(".log"):
        return file_bytes.decode("utf-8", errors="replace")

    if name.endswith(".csv"):
        text = file_bytes.decode("utf-8", errors="replace")
        rows = csv.reader(io.StringIO(text))

        lines = []

        for row in rows:
            lines.append(" | ".join(str(cell) for cell in row))

        return "\n".join(lines)

    if name.endswith(".json"):
        text = file_bytes.decode("utf-8", errors="replace")

        try:
            data = json.loads(text)
            return json.dumps(
                data,
                ensure_ascii=False,
                indent=2
            )
        except json.JSONDecodeError:
            return text

    if name.endswith(".pdf"):
        pdf_file = io.BytesIO(file_bytes)
        reader = PdfReader(pdf_file)

        pages = []

        for page in reader.pages:
            page_text = page.extract_text() or ""
            pages.append(page_text)

        return "\n\n".join(pages)

    if name.endswith(".docx"):
        document = Document(io.BytesIO(file_bytes))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        return "\n".join(paragraphs)

    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        workbook = load_workbook(
            io.BytesIO(file_bytes),
            read_only=True,
            data_only=True
        )

        lines = []

        for worksheet in workbook.worksheets:

            lines.append(
                f"[Sheet: {worksheet.title}]"
            )

            for row in worksheet.iter_rows(
                values_only=True
            ):
                values = []

                for value in row:
                    if value is None:
                        values.append("")
                    else:
                        values.append(str(value))

                lines.append(
                    " | ".join(values)
                )

        return "\n".join(lines)

    raise ValueError(
        "Unsupported file type. "
        "Supported: PDF, DOCX, XLSX, XLSM, TXT, MD, CSV, JSON, LOG."
    )


@app.post("/api/file/upload")
async def upload_file(
    file: UploadFile = File(...),
    http_request: Request = None,
):
    authorization = (
        http_request.headers
        .get("Authorization", "")
        .strip()
        if http_request
        else ""
    )

    if not authorization.startswith("Bearer "):
        raise HTTPException(
            status_code=401,
            detail="Authentication required."
        )

    identity_token = authorization[7:].strip()

    if not identity_token:
        raise HTTPException(
            status_code=401,
            detail="Authentication required."
        )

    try:
        authenticated_user_id = _verify_identity(
            identity_token
        )
    except ValueError:
        raise HTTPException(
            status_code=401,
            detail="Invalid or expired identity token."
        )

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No file selected."
        )

    allowed_extensions = {
        ".pdf",
        ".docx",
        ".xlsx",
        ".xlsm",
        ".txt",
        ".md",
        ".csv",
        ".json",
        ".log",
    }

    extension = os.path.splitext(
        file.filename
    )[1].lower()

    if extension not in allowed_extensions:
        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported file type. "
                "Supported: PDF, DOCX, XLSX, XLSM, "
                "TXT, MD, CSV, JSON, LOG."
            )
        )

    file_bytes = await file.read()

    if len(file_bytes) > MAX_UPLOAD_SIZE:
        raise HTTPException(
            status_code=413,
            detail="File is too large. Maximum size is 10 MB."
        )

    try:
        extracted_text = extract_text_from_file(
            file.filename,
            file_bytes
        )
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Unable to read file: {type(exc).__name__}"
        )

    extracted_text = extracted_text.strip()

    if not extracted_text:
        raise HTTPException(
            status_code=400,
            detail="No readable text was found in this file."
        )

    if len(extracted_text) > MAX_EXTRACTED_TEXT:
        extracted_text = extracted_text[
            :MAX_EXTRACTED_TEXT
        ]

    return {
        "name": "ISMAIL AI",
        "filename": file.filename,
        "content_type": file.content_type,
        "size": len(file_bytes),
        "text": extracted_text,
        "user_id": authenticated_user_id,
    }





class ChatRequest(BaseModel):
    message: str = Field(..., max_length=12000)
    user_id: str = Field("", max_length=200, pattern=r"^[A-Za-z0-9._:-]*$")
    chat_id: str = Field("", max_length=200)
    file_context: str = Field("", max_length=65000)


class ChatHistoryMessage(BaseModel):
    chat_id: str = Field(..., min_length=1, max_length=200)
    role: str = Field(..., pattern=r"^(user|assistant)$")
    message: str = Field(..., min_length=1, max_length=12000)


class ChatHistoryClearRequest(BaseModel):
    pass


class ResearchRequest(BaseModel):
    question: str = Field(..., min_length=1, max_length=12000)
    max_sources: int = Field(5, ge=1, le=10)


class KnowledgeSaveRequest(BaseModel):
    question: str
    answer: str
    topic: str = ""
    source: str = ""
    source_url: str = ""
    verified: bool = False
    confidence: str = "medium"
    expires_at: str | None = None
    knowledge_type: str = "permanent"


class KnowledgeLookupRequest(BaseModel):
    question: str
    knowledge_type: str | None = None


@app.get("/")
def root():
    return {
        "name": "ISMAIL AI",
        "status": "online",
        "version": "0.1.0"
    }


@app.get("/api/ai/status")
def ai_status():
    return {
        "name": "ISMAIL AI",
        "engine": ai_engine.status(),
        "build": {
            "git_commit": os.getenv(
                "RENDER_GIT_COMMIT",
                "local",
            ),
            "structured_live_formatter": True,
        },
    }


@app.post("/api/question/understand")
def understand_question(request: ChatRequest):
    try:
        return {
            "name": "ISMAIL AI",
            "question": ai_engine.understand_question(
                request.message
            )
        }
    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail=str(exc)
        )

    
@app.post("/api/research")
def research(
    request: ResearchRequest,
    http_request: Request,
):
    try:
        authorization = (
            http_request.headers
            .get("Authorization", "")
            .strip()
        )

        if not authorization.startswith("Bearer "):
            raise HTTPException(
                status_code=401,
                detail="Authentication required.",
            )

        identity_token = authorization[7:].strip()

        if not identity_token:
            raise HTTPException(
                status_code=401,
                detail="Authentication required.",
            )

        try:
            _verify_identity(identity_token)
        except ValueError:
            raise HTTPException(
                status_code=401,
                detail="Invalid or expired identity token.",
            )

        client_ip = (
            http_request.client.host
            if http_request.client
            else "unknown"
        )

        if not research_rate_limiter.check(client_ip):
            raise HTTPException(
                status_code=429,
                detail=(
                    "Too many research requests. "
                    "Please try again later."
                ),
            )


        if request.max_sources < 2:
            raise ValueError(
                "max_sources must be at least 2."
            )
        if request.max_sources > 10:
            raise ValueError(
                "max_sources cannot be greater than 10."
            )
        question_info = ai_engine.understand_question(
            request.question
        )
        evidence = web_research.collect_evidence(
            request.question,
            max_sources=request.max_sources,
        )
        return {
            "name": "ISMAIL AI",
            "question": question_info,
            "research": evidence,
        }
    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail=str(exc)
        )
    except RuntimeError as exc:
        raise HTTPException(
            status_code=502,
            detail=str(exc)
        )
@app.post("/api/knowledge/save")
def save_knowledge(
    request: KnowledgeSaveRequest,
    http_request: Request,
):
    authorization = (
        http_request.headers
        .get("Authorization", "")
        .strip()
    )

    if not authorization.startswith("Bearer "):
        raise HTTPException(
            status_code=401,
            detail="Authentication required.",
        )

    identity_token = authorization[7:].strip()

    if not identity_token:
        raise HTTPException(
            status_code=401,
            detail="Authentication required.",
        )

    try:
        _verify_identity(identity_token)
    except ValueError:
        raise HTTPException(
            status_code=401,
            detail="Invalid or expired identity token.",
        )

    admin_key = (
        http_request.headers
        .get("X-Admin-Key", "")
        .strip()
    )

    if (
        not admin_key
        or not hmac.compare_digest(
            admin_key,
            ADMIN_KNOWLEDGE_KEY,
        )
    ):
        raise HTTPException(
            status_code=403,
            detail="Knowledge administration access denied.",
        )

    try:
        knowledge_id = knowledge_base.save(
            question=request.question,
            answer=request.answer,
            topic=request.topic,
            source=request.source,
            source_url=request.source_url,
            verified=request.verified,
            confidence=request.confidence,
            expires_at=request.expires_at,
            knowledge_type=request.knowledge_type,
        )

        return {
            "name": "ISMAIL AI",
            "saved": True,
            "id": knowledge_id,
        }

    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail=str(exc),
        )


    
@app.post("/api/knowledge/lookup")


def lookup_knowledge(
    request: KnowledgeLookupRequest,
    http_request: Request,
):
    try:
        authorization = (
            http_request.headers
            .get("Authorization", "")
            .strip()
        )

        if not authorization.startswith("Bearer "):
            raise HTTPException(
                status_code=401,
                detail="Authentication required.",
            )

        identity_token = authorization[7:].strip()

        if not identity_token:
            raise HTTPException(
                status_code=401,
                detail="Authentication required.",
            )

        try:
            _verify_identity(identity_token)
        except ValueError:
            raise HTTPException(
                status_code=401,
                detail="Invalid or expired identity token.",
            )

        result = knowledge_base.get(
            question=request.question,
            knowledge_type=request.knowledge_type,
        )

        return {
            "name": "ISMAIL AI",
            "found": result is not None,
            "knowledge": result,
        }

    except HTTPException:
        raise

    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail=str(exc),
        )



@app.get("/api/chat/history")
def get_chat_history(
    http_request: Request,
    chat_id: Optional[str] = None,
):
    try:
        authorization = http_request.headers.get("Authorization", "").strip()

        if not authorization.startswith("Bearer "):
            raise HTTPException(
                status_code=401,
                detail="Authentication required.",
            )

        identity_token = authorization[7:].strip()

        if not identity_token:
            raise HTTPException(
                status_code=401,
                detail="Authentication required.",
            )

        try:
            authenticated_user_id = _verify_identity(identity_token)
        except ValueError:
            raise HTTPException(
                status_code=401,
                detail="Invalid or expired identity token.",
            )

        if not chat_id:
            return {
                "name": "ISMAIL AI",
                "user_id": authenticated_user_id,
                "history": [],
            }

        return {
            "name": "ISMAIL AI",
            "user_id": authenticated_user_id,
            "history": knowledge_base.get_chat_history(
                authenticated_user_id,
                chat_id,
            ),
        }

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail="Chat history operation failed. Please try again later."
        )


@app.post("/api/chat/history")
def save_chat_history(
    request: ChatHistoryMessage,
    http_request: Request,
):
    try:
        authorization = http_request.headers.get("Authorization", "").strip()

        if not authorization.startswith("Bearer "):
            raise HTTPException(
                status_code=401,
                detail="Authentication required.",
            )

        identity_token = authorization[7:].strip()

        if not identity_token:
            raise HTTPException(
                status_code=401,
                detail="Authentication required.",
            )

        try:
            authenticated_user_id = _verify_identity(identity_token)
        except ValueError:
            raise HTTPException(
                status_code=401,
                detail="Invalid or expired identity token.",
            )

        saved = knowledge_base.save_chat_message(
            authenticated_user_id,
            request.chat_id,
            request.role,
            request.message,
        )

        if not saved:
            raise HTTPException(
                status_code=400,
                detail="Unable to save chat message.",
            )

        return {
            "name": "ISMAIL AI",
            "saved": True,
            "user_id": authenticated_user_id,
        }

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail="Chat history operation failed. Please try again later."
        )


@app.delete("/api/chat/history")
def clear_chat_history(http_request: Request):
    try:
        authorization = http_request.headers.get("Authorization", "").strip()

        if not authorization.startswith("Bearer "):
            raise HTTPException(
                status_code=401,
                detail="Authentication required.",
            )

        identity_token = authorization[7:].strip()

        if not identity_token:
            raise HTTPException(
                status_code=401,
                detail="Authentication required.",
            )

        try:
            authenticated_user_id = _verify_identity(identity_token)
        except ValueError:
            raise HTTPException(
                status_code=401,
                detail="Invalid or expired identity token.",
            )

        cleared = knowledge_base.clear_chat_history(
            authenticated_user_id
        )

        if not cleared:
            raise HTTPException(
                status_code=400,
                detail="Unable to clear chat history.",
            )

        return {
            "name": "ISMAIL AI",
            "cleared": True,
            "user_id": authenticated_user_id,
        }

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail="Chat history operation failed. Please try again later."
        )




    

    
@app.post("/api/chat")
async def chat(
    request: ChatRequest,
    http_request: Request,
    credentials: HTTPAuthorizationCredentials = Depends(security),
):
    try:
        identity_token = credentials.credentials

        try:
            authenticated_user_id = _verify_identity(identity_token)
        except ValueError:
            raise HTTPException(
                status_code=401,
                detail="Invalid or expired identity token.",
            )

        client_ip = (
            http_request.client.host
            if http_request.client
            else "unknown"
        )

        if not chat_rate_limiter.check(client_ip):
            raise HTTPException(
                status_code=429,
                detail="Too many requests. Please try again later.",
            )

        security = input_security.scan(request.message)

        if not security.allowed:
            raise HTTPException(
                status_code=400,
                detail=security.reason,
            )


        question = ai_engine.understand_question(
            request.message
        )


        response = ai_engine.generate(
            request.message,
            authenticated_user_id,
            request.file_context,
            request.chat_id,
        )

        action = None

        try:
            parsed_response = json.loads(response)

            if isinstance(parsed_response, dict):
                if parsed_response.get("action") != "none":
                    action = parsed_response

        except (json.JSONDecodeError, TypeError):
            pass

        return {
            "name": "ISMAIL AI",
            "question": question,
            "response": response if action is None else "",
            "action": action
        }

    except HTTPException:
        raise

    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail=str(exc)
        )

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail="AI generation failed. Please try again later."
        )




@app.websocket("/ws/voice")
async def voice_websocket(websocket: WebSocket):
    """
    ISMAIL AI realtime voice transport.
    Flow:
        audio_start
        -> continuous audio chunks
        -> audio_end
        -> STT
    """
    await websocket.accept()
    try:
        auth_message = await websocket.receive_json()
        if not isinstance(auth_message, dict):
            await websocket.close(code=1008)
            return
        if auth_message.get("type") != "auth":
            await websocket.close(code=1008)
            return
        identity_token = auth_message.get("token", "")
        try:
            authenticated_user_id = _verify_identity(identity_token)
        except (ValueError, TypeError):
            await websocket.send_json({
                "type": "error",
                "code": "AUTH_FAILED",
                "message": "Invalid or expired identity token.",
            })
            await websocket.close(code=1008)
            return
        chat_id = str(auth_message.get("chat_id", "")).strip()
        if not chat_id:
            await websocket.send_json({
                "type": "error",
                "code": "CHAT_ID_REQUIRED",
                "message": "chat_id is required.",
            })
            await websocket.close(code=1008)
            return
        websocket._ismail_audio_buffer = bytearray()
        websocket._ismail_sample_rate = 16000
        websocket._ismail_voice_active = False
        await websocket.send_json({
            "type": "auth_ok",
            "user_id": authenticated_user_id,
            "chat_id": chat_id,
            "status": "ready",
        })
        while True:
            message = await websocket.receive_json()
            if not isinstance(message, dict):
                await websocket.send_json({
                    "type": "error",
                    "code": "INVALID_MESSAGE",
                    "message": "Message must be a JSON object.",
                })
                continue
            message_type = message.get("type")
            if message_type == "ping":
                await websocket.send_json({
                    "type": "pong",
                })
            elif message_type == "close":
                await websocket.close(code=1000)
                return
            elif message_type == "audio_start":
                sample_rate = message.get("sample_rate", 16000)
                channels = message.get("channels", 1)
                if (
                    not isinstance(sample_rate, int)
                    or isinstance(sample_rate, bool)
                    or sample_rate <= 0
                ):
                    await websocket.send_json({
                        "type": "error",
                        "code": "INVALID_SAMPLE_RATE",
                        "message": "sample_rate must be a positive integer.",
                    })
                    continue
                if (
                    not isinstance(channels, int)
                    or isinstance(channels, bool)
                    or channels <= 0
                ):
                    await websocket.send_json({
                        "type": "error",
                        "code": "INVALID_CHANNELS",
                        "message": "channels must be a positive integer.",
                    })
                    continue
                websocket._ismail_audio_buffer = bytearray()
                websocket._ismail_sample_rate = sample_rate
                websocket._ismail_channels = channels
                websocket._ismail_voice_active = True
                await websocket.send_json({
                    "type": "audio_started",
                    "format": "pcm16",
                    "sample_rate": sample_rate,
                    "channels": channels,
                    "status": "listening",
                })
            elif message_type == "audio":
                audio_format = str(
                    message.get("format", "")
                ).strip().lower()
                sample_rate = message.get(
                    "sample_rate",
                    getattr(websocket, "_ismail_sample_rate", 16000),
                )
                audio_data = message.get("data", "")
                if not getattr(
                    websocket,
                    "_ismail_voice_active",
                    False,
                ):
                    await websocket.send_json({
                        "type": "error",
                        "code": "AUDIO_SESSION_NOT_STARTED",
                        "message": "Send audio_start before audio.",
                    })
                    continue
                if audio_format not in {"pcm16", "audio/pcm"}:
                    await websocket.send_json({
                        "type": "error",
                        "code": "UNSUPPORTED_AUDIO_FORMAT",
                        "message": "Supported audio format: pcm16.",
                    })
                    continue
                if (
                    not isinstance(sample_rate, int)
                    or isinstance(sample_rate, bool)
                    or sample_rate <= 0
                ):
                    await websocket.send_json({
                        "type": "error",
                        "code": "INVALID_SAMPLE_RATE",
                        "message": "sample_rate must be a positive integer.",
                    })
                    continue
                if not isinstance(audio_data, str) or not audio_data:
                    await websocket.send_json({
                        "type": "error",
                        "code": "INVALID_AUDIO_DATA",
                        "message": "audio data must be a base64 string.",
                    })
                    continue
                try:
                    pcm_bytes = base64.b64decode(
                        audio_data,
                        validate=True,
                    )
                except (ValueError, TypeError):
                    await websocket.send_json({
                        "type": "error",
                        "code": "INVALID_BASE64",
                        "message": "audio data is not valid base64.",
                    })
                    continue
                if not pcm_bytes:
                    continue
                if len(pcm_bytes) % 2 != 0:
                    await websocket.send_json({
                        "type": "error",
                        "code": "INVALID_PCM16",
                        "message": (
                            "PCM16 audio must contain "
                            "an even number of bytes."
                        ),
                    })
                    continue
                audio_buffer = getattr(
                    websocket,
                    "_ismail_audio_buffer",
                    bytearray(),
                )
                audio_buffer.extend(pcm_bytes)
                websocket._ismail_audio_buffer = audio_buffer
                await websocket.send_json({
                    "type": "audio_ack",
                    "format": "pcm16",
                    "sample_rate": sample_rate,
                    "bytes_received": len(pcm_bytes),
                    "buffered_bytes": len(audio_buffer),
                    "samples_received": len(pcm_bytes) // 2,
                    "status": "accepted",
                })
            elif message_type == "audio_end":
                audio_buffer = getattr(
                    websocket,
                    "_ismail_audio_buffer",
                    bytearray(),
                )
                sample_rate = getattr(
                    websocket,
                    "_ismail_sample_rate",
                    16000,
                )
                channels = getattr(
                    websocket,
                    "_ismail_channels",
                    1,
                )
                websocket._ismail_voice_active = False
                pcm_bytes = bytes(audio_buffer)
                if not pcm_bytes:
                    await websocket.send_json({
                        "type": "audio_committed",
                        "bytes": 0,
                        "samples": 0,
                        "status": "empty",
                    })
                    continue
                from backend.voice.speech_to_text import speech_to_text
                transcript = speech_to_text.transcribe_pcm16(
                    pcm_bytes,
                    sample_rate=sample_rate,
                    channels=channels,
                )
                await websocket.send_json({
                    "type": "audio_committed",
                    "bytes": len(pcm_bytes),
                    "samples": len(pcm_bytes) // 2,
                    "status": "committed",
                })
                if transcript:
                    await websocket.send_json({
                        "type": "transcript",
                        "text": transcript,
                        "final": True,
                    })
                else:
                    await websocket.send_json({
                        "type": "transcript",
                        "text": "",
                        "final": True,
                        "status": "stt_provider_not_configured",
                    })
                websocket._ismail_audio_buffer = bytearray()
            else:
                await websocket.send_json({
                    "type": "error",
                    "code": "UNSUPPORTED_MESSAGE",
                    "message": "Unsupported realtime message type.",
                })
    except WebSocketDisconnect:
        return
    except Exception:
        try:
            await websocket.close(code=1011)
        except Exception:
            pass

if __name__ == "__main__":
    import uvicorn

    uvicorn.run(
        "backend.server:app",
        host="0.0.0.0",
        port=8000,
        reload=True
    )
















