import csv
import io
import json

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_EXTRACTED_CHARS = 60000


ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".xlsm",
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".log",
}


def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
    """
    Extract readable text from supported document formats.
    """

    if not filename:
        raise ValueError("Filename is required.")

    if len(file_bytes) > MAX_FILE_SIZE:
        raise ValueError("File is too large. Maximum size is 10 MB.")

    extension = "." + filename.lower().split(".")[-1]

    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(
            "Unsupported file type. "
            "Supported: PDF, DOCX, XLSX, XLSM, TXT, MD, CSV, JSON, LOG."
        )

    if extension == ".pdf":
        return _extract_pdf(file_bytes)

    if extension == ".docx":
        return _extract_docx(file_bytes)

    if extension in {".xlsx", ".xlsm"}:
        return _extract_excel(file_bytes)

    if extension in {".txt", ".md", ".log"}:
        return _extract_text(file_bytes)

    if extension == ".csv":
        return _extract_csv(file_bytes)

    if extension == ".json":
        return _extract_json(file_bytes)

    raise ValueError("Could not process this file.")


def _limit_text(text: str) -> str:
    text = text.strip()

    if len(text) > MAX_EXTRACTED_CHARS:
        text = (
            text[:MAX_EXTRACTED_CHARS]
            + "\n\n[Document content truncated at 60,000 characters.]"
        )

    return text


def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))

    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""

        if text.strip():
            pages.append(
                f"--- Page {page_number} ---\n{text.strip()}"
            )

    return _limit_text("\n\n".join(pages))


def _extract_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))

    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table_number, table in enumerate(document.tables, start=1):
        parts.append(f"--- Table {table_number} ---")

        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(values))

    return _limit_text("\n".join(parts))


def _extract_excel(file_bytes: bytes) -> str:
    workbook = load_workbook(
        io.BytesIO(file_bytes),
        read_only=True,
        data_only=True,
    )

    parts = []

    for worksheet in workbook.worksheets:
        parts.append(f"--- Sheet: {worksheet.title} ---")

        for row in worksheet.iter_rows(values_only=True):
            values = []

            for value in row:
                if value is None:
                    values.append("")
                else:
                    values.append(str(value))

            if any(value.strip() for value in values):
                parts.append(" | ".join(values))

    return _limit_text("\n".join(parts))


def _extract_text(file_bytes: bytes) -> str:
    text = file_bytes.decode("utf-8", errors="replace")
    return _limit_text(text)


def _extract_csv(file_bytes: bytes) -> str:
    text = file_bytes.decode("utf-8-sig", errors="replace")

    reader = csv.reader(io.StringIO(text))

    rows = []

    for row in reader:
        rows.append(" | ".join(row))

    return _limit_text("\n".join(rows))


def _extract_json(file_bytes: bytes) -> str:
    text = file_bytes.decode("utf-8", errors="replace")

    data = json.loads(text)

    formatted = json.dumps(
        data,
        ensure_ascii=False,
        indent=2,
    )

    return _limit_text(formatted)